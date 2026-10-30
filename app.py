import streamlit as st
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO

# ==================== DEFAULTS ====================
DEFAULT_BASE_RATE = 90.0

# Default percentages (extra charge on top of base). 
# 0 = Normal Charge, 25 = +25%, 50 = +50%, etc.
DEFAULT_PERCENTAGES = {
    "Weekday": {
        "normal": 0,      # Normal Charge
        "evening": 25,
        "night": 50
    },
    "Saturday": {
        "normal": 25,
        "evening": 50,
        "night": 100
    },
    "Sunday": {
        "normal": 50,
        "evening": 100,
        "night": 150
    },
    "Public Holiday": {
        "normal": 100,
        "evening": 150,
        "night": 200
    }
}

DAY_LABELS = {
    "Weekday": "Isnin – Jumaat",
    "Saturday": "Sabtu",
    "Sunday": "Ahad",
    "Public Holiday": "Cuti Umum"
}

SHIFT_LABELS = {
    "normal": "08:00 – 17:00",
    "evening": "17:00 – 24:00",
    "night": "24:00 – 08:00"
}

# ==================== PAGE SETUP ====================
st.set_page_config(
    page_title="Manpower Rate Calculator",
    page_icon="👷",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
    .main-title {
        font-size: 1.9rem;
        font-weight: 700;
        color: #1e3a5f;
        margin-bottom: 0.1rem;
    }
    .sub-title {
        color: #64748b;
        font-size: 0.95rem;
        margin-bottom: 1.5rem;
    }
    div[data-testid="stMetricValue"] {
        font-size: 1.3rem;
    }
</style>
""", unsafe_allow_html=True)

# ==================== SESSION STATE ====================
if "base_rate" not in st.session_state:
    st.session_state.base_rate = DEFAULT_BASE_RATE
if "hours_normal" not in st.session_state:
    st.session_state.hours_normal = 8.0
if "hours_evening" not in st.session_state:
    st.session_state.hours_evening = 2.0
if "hours_night" not in st.session_state:
    st.session_state.hours_night = 1.0
if "percentages" not in st.session_state:
    st.session_state.percentages = {
        day: shifts.copy() for day, shifts in DEFAULT_PERCENTAGES.items()
    }

# ==================== HELPERS ====================
def pct_to_multiplier(pct: float) -> float:
    """Convert percentage extra charge to multiplier. 0% → 1.0, 25% → 1.25, 100% → 2.0"""
    return round(1.0 + (pct / 100.0), 4)

def get_rate(day_type: str, shift: str, base: float, percentages: dict) -> float:
    pct = percentages[day_type][shift]
    mult = pct_to_multiplier(pct)
    return round(base * mult, 2)

def calculate_all(num_days: dict, hours: dict, base: float, percentages: dict):
    results = []
    grand_total_hours = 0.0
    grand_total_cost = 0.0

    for day_type, count in num_days.items():
        if count <= 0:
            continue

        rate_n = get_rate(day_type, "normal", base, percentages)
        rate_e = get_rate(day_type, "evening", base, percentages)
        rate_ni = get_rate(day_type, "night", base, percentages)

        total_h_n = count * hours["normal"]
        total_h_e = count * hours["evening"]
        total_h_ni = count * hours["night"]

        cost_n = round(total_h_n * rate_n, 2)
        cost_e = round(total_h_e * rate_e, 2)
        cost_ni = round(total_h_ni * rate_ni, 2)
        total_cost = cost_n + cost_e + cost_ni
        total_hours = total_h_n + total_h_e + total_h_ni

        results.append({
            "day_type": day_type,
            "label": DAY_LABELS[day_type],
            "days": count,
            "hours_normal": total_h_n,
            "hours_evening": total_h_e,
            "hours_night": total_h_ni,
            "total_hours": total_hours,
            "rate_normal": rate_n,
            "rate_evening": rate_e,
            "rate_night": rate_ni,
            "pct_normal": percentages[day_type]["normal"],
            "pct_evening": percentages[day_type]["evening"],
            "pct_night": percentages[day_type]["night"],
            "cost_normal": cost_n,
            "cost_evening": cost_e,
            "cost_night": cost_ni,
            "total_cost": total_cost
        })

        grand_total_hours += total_hours
        grand_total_cost += total_cost

    return results, grand_total_hours, grand_total_cost

def export_excel(results, grand_hours, grand_cost, base, hours, num_days, percentages):
    wb = Workbook()
    ws = wb.active
    ws.title = "Manpower Calculation"

    header_fill = PatternFill(start_color="1e3a5f", end_color="1e3a5f", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    money_fill = PatternFill(start_color="dcfce7", end_color="dcfce7", fill_type="solid")
    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    ws.merge_cells('A1:I1')
    ws['A1'] = "MANPOWER RATE CALCULATOR - PPM TA EXPERT"
    ws['A1'].font = Font(bold=True, size=14, color="1e3a5f")

    ws['A2'] = f"Base Manhour Rate: RM {base:.2f} / jam"
    ws['A3'] = f"Standard Hours/Day → Normal: {hours['normal']}h | Evening: {hours['evening']}h | Night: {hours['night']}h"
    ws['A4'] = f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}"

    # Charge Rate table used
    ws['A6'] = "CHARGE RATE YANG DIGUNAKAN (%)"
    ws['A6'].font = Font(bold=True, size=12)

    headers_rate = ["Jenis Hari", "08:00-17:00", "17:00-24:00", "24:00-08:00"]
    for col, h in enumerate(headers_rate, 1):
        cell = ws.cell(row=7, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin

    row = 8
    for dt in ["Weekday", "Saturday", "Sunday", "Public Holiday"]:
        ws.cell(row=row, column=1, value=DAY_LABELS[dt]).border = thin
        for col, shift in enumerate(["normal", "evening", "night"], 2):
            pct = percentages[dt][shift]
            cell = ws.cell(row=row, column=col, value=f"{pct}%")
            cell.border = thin
            cell.alignment = Alignment(horizontal="center")
        row += 1

    # Input days
    row += 1
    ws.cell(row=row, column=1, value="BILANGAN HARI").font = Font(bold=True, size=12)
    row += 1
    ws.cell(row=row, column=1, value="Jenis Hari").font = header_font
    ws.cell(row=row, column=1).fill = header_fill
    ws.cell(row=row, column=2, value="Bilangan").font = header_font
    ws.cell(row=row, column=2).fill = header_fill
    row += 1
    for dt, label in DAY_LABELS.items():
        ws.cell(row=row, column=1, value=label).border = thin
        ws.cell(row=row, column=2, value=num_days.get(dt, 0)).border = thin
        row += 1

    # Calculation details
    row += 1
    ws.cell(row=row, column=1, value="PERINCIAN KIRAAN").font = Font(bold=True, size=12)
    row += 1

    headers = ["Jenis Hari", "Hari", "Jam Normal", "Jam Petang", "Jam Malam",
               "Kos Normal (RM)", "Kos Petang (RM)", "Kos Malam (RM)", "Jumlah Kos (RM)"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row += 1
    for r in results:
        ws.cell(row=row, column=1, value=r["label"]).border = thin
        ws.cell(row=row, column=2, value=r["days"]).border = thin
        ws.cell(row=row, column=3, value=r["hours_normal"]).border = thin
        ws.cell(row=row, column=4, value=r["hours_evening"]).border = thin
        ws.cell(row=row, column=5, value=r["hours_night"]).border = thin

        for col, key in enumerate(["cost_normal", "cost_evening", "cost_night", "total_cost"], 6):
            cell = ws.cell(row=row, column=col, value=r[key])
            cell.border = thin
            cell.number_format = '#,##0.00'
            if key == "total_cost":
                cell.fill = money_fill
                cell.font = Font(bold=True)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="GRAND TOTAL").font = Font(bold=True, size=12)
    ws.cell(row=row, column=3, value=grand_hours).font = Font(bold=True)
    cell = ws.cell(row=row, column=9, value=round(grand_cost, 2))
    cell.font = Font(bold=True, size=13, color="166534")
    cell.fill = money_fill
    cell.number_format = '"RM "#,##0.00'

    ws.column_dimensions['A'].width = 22
    for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
        ws.column_dimensions[col].width = 14

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

# ==================== SIDEBAR ====================
with st.sidebar:
    st.markdown("### ⚙️ Tetapan Asas")

    base_rate = st.number_input(
        "Base Manhour Rate (RM / jam)",
        min_value=1.0,
        max_value=2000.0,
        value=st.session_state.base_rate,
        step=0.5,
        format="%.2f"
    )
    st.session_state.base_rate = base_rate

    st.markdown("---")
    st.markdown("### 🕐 Standard Jam Kerja (sehari)")
    st.caption("Diguna untuk semua jenis hari")

    h_normal = st.number_input("Jam Normal (08:00–17:00)", min_value=0.0, max_value=24.0,
                               value=st.session_state.hours_normal, step=0.5)
    h_evening = st.number_input("Jam Petang (17:00–24:00)", min_value=0.0, max_value=24.0,
                                value=st.session_state.hours_evening, step=0.5)
    h_night = st.number_input("Jam Malam (24:00–08:00)", min_value=0.0, max_value=24.0,
                              value=st.session_state.hours_night, step=0.5)

    st.session_state.hours_normal = h_normal
    st.session_state.hours_evening = h_evening
    st.session_state.hours_night = h_night

    st.markdown("---")
    if st.button("🔄 Reset Percentage ke Default", use_container_width=True):
        st.session_state.percentages = {
            day: shifts.copy() for day, shifts in DEFAULT_PERCENTAGES.items()
        }
        st.rerun()

# ==================== MAIN ====================
st.markdown('<p class="main-title">👷 Manpower Rate Calculator</p>', unsafe_allow_html=True)
st.markdown('<p class="sub-title">PPM TA Expert • Custom Charge Rate + Auto kira bilangan hari</p>', unsafe_allow_html=True)

# ========== EDITABLE CHARGE RATE TABLE ==========
st.markdown("### 📝 Charge Rate (Percentage) — Boleh Edit")
st.caption("Masukkan % tambahan di atas Base Rate. Contoh: 0 = Normal Charge, 25 = +25% (1.25×), 100 = +100% (2.0×)")

# Create editable grid
pct = st.session_state.percentages

# Header row
cols = st.columns([2.2, 1.5, 1.5, 1.5])
cols[0].markdown("**Jenis Hari**")
cols[1].markdown("**08:00 – 17:00**")
cols[2].markdown("**17:00 – 24:00**")
cols[3].markdown("**24:00 – 08:00**")

# Data rows
for day_key in ["Weekday", "Saturday", "Sunday", "Public Holiday"]:
    cols = st.columns([2.2, 1.5, 1.5, 1.5])
    cols[0].write(DAY_LABELS[day_key])

    for i, shift in enumerate(["normal", "evening", "night"]):
        new_val = cols[i + 1].number_input(
            label=f"{day_key}_{shift}",
            min_value=0,
            max_value=500,
            value=int(pct[day_key][shift]),
            step=5,
            key=f"pct_{day_key}_{shift}",
            label_visibility="collapsed"
        )
        st.session_state.percentages[day_key][shift] = new_val

# Live rate preview
with st.expander("👀 Preview Rate (RM) berdasarkan percentage di atas", expanded=False):
    preview_data = []
    for day_key in ["Weekday", "Saturday", "Sunday", "Public Holiday"]:
        row = {"Jenis Hari": DAY_LABELS[day_key]}
        for shift, label in SHIFT_LABELS.items():
            rate = get_rate(day_key, shift, base_rate, st.session_state.percentages)
            pct_val = st.session_state.percentages[day_key][shift]
            row[label] = f"RM {rate:.2f}  ({pct_val}%)"
        preview_data.append(row)
    st.dataframe(pd.DataFrame(preview_data), use_container_width=True, hide_index=True)

st.markdown("---")

# ========== DAYS INPUT ==========
st.markdown("### 📅 Masukkan Bilangan Hari")

c1, c2, c3, c4 = st.columns(4)

with c1:
    num_weekday = st.number_input("Isnin – Jumaat", min_value=0, max_value=365, value=0, step=1)
with c2:
    num_saturday = st.number_input("Sabtu", min_value=0, max_value=100, value=0, step=1)
with c3:
    num_sunday = st.number_input("Ahad", min_value=0, max_value=100, value=0, step=1)
with c4:
    num_ph = st.number_input("Cuti Umum", min_value=0, max_value=50, value=0, step=1)

num_days = {
    "Weekday": num_weekday,
    "Saturday": num_saturday,
    "Sunday": num_sunday,
    "Public Holiday": num_ph
}

hours = {
    "normal": h_normal,
    "evening": h_evening,
    "night": h_night
}

total_days = sum(num_days.values())

# ========== RESULTS ==========
if total_days > 0:
    results, grand_hours, grand_cost = calculate_all(
        num_days, hours, base_rate, st.session_state.percentages
    )

    st.markdown("---")
    st.markdown("### 📊 Keputusan")

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Jumlah Hari", f"{total_days} hari")
    m2.metric("Jumlah Jam", f"{grand_hours:,.1f} jam")
    m3.metric("Base Rate", f"RM {base_rate:.2f}")
    m4.metric("GRAND TOTAL", f"RM {grand_cost:,.2f}")

    st.markdown("")
    st.markdown("#### Perincian Mengikut Jenis Hari")

    df = pd.DataFrame([
        {
            "Jenis Hari": r["label"],
            "Hari": r["days"],
            "Jam Normal": r["hours_normal"],
            "Jam Petang": r["hours_evening"],
            "Jam Malam": r["hours_night"],
            "Jumlah Jam": r["total_hours"],
            "Kos Normal (RM)": r["cost_normal"],
            "Kos Petang (RM)": r["cost_evening"],
            "Kos Malam (RM)": r["cost_night"],
            "Jumlah Kos (RM)": r["total_cost"]
        }
        for r in results
    ])

    df_display = df.copy()
    for col in ["Kos Normal (RM)", "Kos Petang (RM)", "Kos Malam (RM)", "Jumlah Kos (RM)"]:
        df_display[col] = df_display[col].apply(lambda x: f"{x:,.2f}")

    st.dataframe(df_display, use_container_width=True, hide_index=True)

    # Summary
    st.markdown("#### Ringkasan Kos")
    col_a, col_b = st.columns(2)

    with col_a:
        st.write("**Mengikut Shift**")
        st.write(f"• Normal: **RM {sum(r['cost_normal'] for r in results):,.2f}**")
        st.write(f"• Petang: **RM {sum(r['cost_evening'] for r in results):,.2f}**")
        st.write(f"• Malam: **RM {sum(r['cost_night'] for r in results):,.2f}**")

    with col_b:
        st.write("**Mengikut Jenis Hari**")
        for r in results:
            st.write(f"• {r['label']}: **RM {r['total_cost']:,.2f}**")

    st.markdown("---")
    excel_data = export_excel(
        results, grand_hours, grand_cost, base_rate, hours, num_days, st.session_state.percentages
    )

    st.download_button(
        label="📄 Download Excel Report",
        data=excel_data,
        file_name=f"Manpower_Calc_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True
    )

else:
    st.info("👆 Masukkan bilangan hari di atas untuk mula kira.")

st.markdown("---")
st.caption("Manpower Rate Calculator • PPM TA Expert • Percentage charge rate boleh diubah sendiri")
